"""Fill the existing '평균 조회수' column of a user's Excel sheet.

Copies the input workbook, finds each sheet's '링크' and '평균 조회수' columns
on the header row (default row 4), then for each data row scrapes the linked
Instagram account and writes the average play count (앱 표기 "조회수") of the
most recent N video posts. Secondary metrics (시청수, 분석 영상 수, 수집 게시물 수)
are attached as a cell comment so the user's column layout is unchanged.

    python execution/fill_views.py input.xlsx --sheets 브랜드 메뉴
"""

from __future__ import annotations

import argparse
import logging
import re
import shutil
import sys
from pathlib import Path

import instagram_scraper
from batch_views import recent_video_metrics
from config import OUTPUT_DIR, RESULTS_LIMIT, ensure_dirs

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("fill_views")

HEADER_ROW = 4
URL_RE = re.compile(r"https?://(?:www\.)?instagram\.com/[^\s\"'<>]+", re.IGNORECASE)


def find_columns(ws, header_row: int = HEADER_ROW) -> tuple[int | None, int | None, int | None]:
    headers: dict[str, int] = {}
    for col, cell in enumerate(ws[header_row], 1):
        if isinstance(cell.value, str):
            headers[cell.value.strip()] = col
    link_col = headers.get("링크") or headers.get("LINK") or headers.get("Link")
    views_col = headers.get("평균 조회수") or headers.get("Avg Views")
    id_col = headers.get("아이디") or headers.get("ID")
    return link_col, views_col, id_col


def extract_url(value) -> str | None:
    if not isinstance(value, str):
        return None
    m = URL_RE.search(value)
    return m.group(0).rstrip(".,);]") if m else None


def fill_sheet(ws, *, limit: int, n_videos: int) -> tuple[int, int]:
    from openpyxl.comments import Comment

    link_col, views_col, _ = find_columns(ws)
    if not link_col or not views_col:
        logger.info("시트 '%s' 건너뜀(필요 컬럼 없음)", ws.title)
        return 0, 0
    logger.info(
        "시트 '%s' 처리 시작 (링크열=%d, 평균 조회수열=%d)",
        ws.title,
        link_col,
        views_col,
    )

    success = miss = 0
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        url = extract_url(ws.cell(row=r, column=link_col).value)
        if not url:
            continue
        try:
            target = instagram_scraper.parse_instagram_url(url)
        except ValueError as exc:
            ws.cell(row=r, column=views_col).comment = Comment(f"잘못된 URL: {exc}", "batch")
            miss += 1
            continue

        handle = target.username or url
        try:
            items = instagram_scraper.scrape(url, limit=limit)
            metrics = recent_video_metrics(items, n=n_videos)
        except Exception as exc:
            logger.error("실패 [row=%d @%s]: %s", r, handle, exc)
            ws.cell(row=r, column=views_col).comment = Comment(f"스크랩 오류: {exc}", "batch")
            miss += 1
            continue

        cell = ws.cell(row=r, column=views_col)
        if metrics["avg_play_count"] is not None:
            cell.value = metrics["avg_play_count"]
            cell.number_format = "#,##0"
            note_lines = [
                f"최근 영상 {metrics['videos_used']}개 평균(재생수)",
                f"시청수 평균: {metrics['avg_view_count']}",
                f"수집 게시물: {metrics['posts_scraped']}건",
            ]
            if metrics["note"]:
                note_lines.insert(0, metrics["note"])
            cell.comment = Comment("\n".join(note_lines), "batch")
            success += 1
            logger.info(
                "  row %d @%s -> 재생수 %s · 시청수 %s · 영상 %d개",
                r,
                handle,
                metrics["avg_play_count"],
                metrics["avg_view_count"],
                metrics["videos_used"],
            )
        else:
            cell.comment = Comment(metrics["note"] or "영상 없음", "batch")
            miss += 1
            logger.info("  row %d @%s -> 영상 없음/미입력", r, handle)

    logger.info("시트 '%s' 완료 (성공 %d · 미입력 %d)", ws.title, success, miss)
    return success, miss


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="엑셀의 기존 '평균 조회수' 컬럼을 채워 넣습니다."
    )
    parser.add_argument("xlsx", help="입력 .xlsx 경로")
    parser.add_argument(
        "--output", default=None, help="출력 .xlsx 경로 (미지정 시 Output/<원본명>_with_views.xlsx)"
    )
    parser.add_argument("--videos", type=int, default=5, help="평균을 낼 최근 영상 수")
    parser.add_argument(
        "--limit", type=int, default=RESULTS_LIMIT, help="계정당 수집할 최대 게시물 수"
    )
    parser.add_argument(
        "--sheets",
        nargs="*",
        default=None,
        help="처리할 시트 이름들 (미지정 시 적용 가능한 모든 시트)",
    )
    args = parser.parse_args(argv)

    from openpyxl import load_workbook

    ensure_dirs()
    in_path = Path(args.xlsx)
    out_path = (
        Path(args.output)
        if args.output
        else OUTPUT_DIR / f"{in_path.stem}_with_views.xlsx"
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(in_path, out_path)
    wb = load_workbook(out_path)

    total_s = total_m = 0
    for ws in wb.worksheets:
        if args.sheets and ws.title not in args.sheets:
            continue
        s, m = fill_sheet(ws, limit=args.limit, n_videos=args.videos)
        total_s += s
        total_m += m

    wb.save(out_path)
    logger.info("✅ 저장 완료: %s | 성공 %d · 미입력 %d", out_path, total_s, total_m)
    return 0


if __name__ == "__main__":
    sys.exit(main())
