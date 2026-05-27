"""Batch view-count report from an Excel sheet of Instagram links.

Reads an .xlsx, auto-detects Instagram URLs in any cell, scrapes each account,
and computes the average view count over its most recent video posts. Results
are written to an .xlsx.

    python execution/batch_views.py input.xlsx
    python execution/batch_views.py input.xlsx --output Output/views.xlsx --videos 5

Note on "조회수": Instagram exposes two metrics on videos/reels —
``videoPlayCount`` (재생수; the large number shown as "조회수" in-app) and
``videoViewCount`` (시청수). Both are reported. Image/Sidecar posts have no
view metric, so only video posts are averaged.
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from datetime import datetime

import instagram_scraper
from config import OUTPUT_DIR, RESULTS_LIMIT, ensure_dirs

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
logger = logging.getLogger("batch_views")

DEFAULT_RECENT_VIDEOS = 5
_URL_RE = re.compile(r"https?://(?:www\.)?instagram\.com/[^\s\"'<>]+", re.IGNORECASE)


def find_instagram_urls(xlsx_path: str) -> list[str]:
    """Return unique, de-duplicated Instagram account/post URLs from any cell."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    seen: dict[str, str] = {}  # dedupe key -> original url
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if not isinstance(cell, str):
                    continue
                for match in _URL_RE.findall(cell):
                    url = match.rstrip(".,);]")
                    try:
                        target = instagram_scraper.parse_instagram_url(url)
                    except ValueError:
                        continue
                    key = target.username or target.shortcode or url
                    seen.setdefault(key.lower(), url)
    wb.close()
    return list(seen.values())


def _as_int(value) -> int | None:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return None
    return n if n >= 0 else None


def _ts_key(item: dict) -> float:
    ts = item.get("timestamp")
    try:
        return datetime.fromisoformat(str(ts).replace("Z", "+00:00")).timestamp()
    except (ValueError, TypeError):
        return 0.0


def _mean(values: list[int]) -> int | None:
    return round(sum(values) / len(values)) if values else None


def recent_video_metrics(items: list[dict], n: int = DEFAULT_RECENT_VIDEOS) -> dict:
    """Average view/play counts over the n most recent video posts."""
    videos = [i for i in items if (i.get("type") or "").lower() == "video"]
    videos.sort(key=_ts_key, reverse=True)
    recent = videos[:n]

    plays = [v for v in (_as_int(i.get("videoPlayCount")) for i in recent) if v is not None]
    views = [v for v in (_as_int(i.get("videoViewCount")) for i in recent) if v is not None]

    note = ""
    if not videos:
        note = "영상 게시물 없음"
    elif len(recent) < n:
        note = f"영상 {len(recent)}개만 분석(요청 {n}개)"

    return {
        "posts_scraped": len(items),
        "videos_found": len(videos),
        "videos_used": len(recent),
        "avg_play_count": _mean(plays),
        "avg_view_count": _mean(views),
        "used_urls": [i.get("url") for i in recent if i.get("url")],
        "note": note,
    }


def process(urls: list[str], limit: int, n: int) -> list[dict]:
    rows = []
    for url in urls:
        logger.info("처리 중: %s", url)
        try:
            items = instagram_scraper.scrape(url, limit=limit)
            metrics = recent_video_metrics(items, n=n)
        except Exception as exc:
            logger.error("실패(%s): %s", url, exc)
            metrics = {
                "posts_scraped": 0,
                "videos_found": 0,
                "videos_used": 0,
                "avg_play_count": None,
                "avg_view_count": None,
                "used_urls": [],
                "note": f"오류: {exc}",
            }
        try:
            handle = instagram_scraper.parse_instagram_url(url).username or ""
        except ValueError:
            handle = ""
        rows.append({"url": url, "handle": handle, **metrics})
    return rows


def write_xlsx(rows: list[dict], out_path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "조회수 분석"
    headers = [
        "입력 링크",
        "계정",
        "평균 조회수(재생수)",
        "평균 시청수",
        "분석 영상 수",
        "수집 게시물 수",
        "비고",
        "분석 대상 게시물",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        ws.append(
            [
                r["url"],
                f"@{r['handle']}" if r["handle"] else "",
                r["avg_play_count"],
                r["avg_view_count"],
                r["videos_used"],
                r["posts_scraped"],
                r["note"],
                "\n".join(r["used_urls"]),
            ]
        )

    widths = [46, 18, 18, 16, 12, 14, 26, 50]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logger.info("결과 저장: %s", out_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="엑셀의 인스타그램 링크별 최근 영상 평균 조회수를 계산합니다."
    )
    parser.add_argument("xlsx", help="인스타그램 링크가 들어 있는 .xlsx 파일")
    parser.add_argument(
        "--output",
        default=str(OUTPUT_DIR / "instagram_views.xlsx"),
        help="출력 .xlsx 경로",
    )
    parser.add_argument(
        "--videos",
        type=int,
        default=DEFAULT_RECENT_VIDEOS,
        help="평균을 낼 최근 영상 수(기본 5)",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=RESULTS_LIMIT,
        help="계정당 수집할 최대 게시물 수(영상 확보용, 기본 30)",
    )
    args = parser.parse_args(argv)

    ensure_dirs()
    urls = find_instagram_urls(args.xlsx)
    if not urls:
        logger.error("엑셀에서 인스타그램 링크를 찾지 못했습니다.")
        return 1
    logger.info("감지된 계정/링크 %d개", len(urls))

    from pathlib import Path

    rows = process(urls, limit=args.limit, n=args.videos)
    write_xlsx(rows, Path(args.output))
    return 0


if __name__ == "__main__":
    sys.exit(main())
